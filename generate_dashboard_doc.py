#!/usr/bin/env python3
"""Generate dashboard descriptions Word document."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT = "/workspace/أوصاف_لوحات_المعلومات.docx"

SECTIONS = [
    {
        "dashboard_ar": "مسح مقاولي الإنشاءات",
        "dashboard_en": "Construction Contractors Survey",
        "tabs": [
            {
                "name_ar": "النتائج الرئيسية لمسح مقاولي الإنشاءات حسب المحافظة",
                "name_en": "Main Results of the Construction Contractors Survey by Governorate",
                "desc_ar": "لوحة تعرض مؤشرات مسح مقاولي الإنشاءات عبر مخططات بيانية وفلاتر للسنة والمحافظة.",
                "desc_en": "A dashboard displaying construction contractors survey indicators through charts and filters for year and governorate.",
                "charts": [],
            }
        ],
    },
    {
        "dashboard_ar": "الطاقة والصناعة",
        "dashboard_en": "Energy and Industry",
        "tabs": [
            {
                "name_ar": "إنتاج ومبيعات المحروقات",
                "name_en": "Fuel Production and Sales",
                "desc_ar": "لوحة تعرض مخططات بيانية لإنتاج ومبيعات المحروقات مع فلاتر للمادة والسنة.",
                "desc_en": "A dashboard displaying charts for fuel production and sales with filters for material and year.",
                "charts": [
                    ("مبيعات", "Sales", "مخطط يعرض مبيعات المادة المحددة حسب السنة.", "A chart showing sales of the selected material by year."),
                    ("إنتاج", "Production", "مخطط يعرض إنتاج المادة المحددة حسب السنة.", "A chart showing production of the selected material by year."),
                ],
            },
            {
                "name_ar": "الإنتاج المحلي من الطاقة الجديدة والمتجددة",
                "name_en": "Local Production of New and Renewable Energy",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لإنتاج الطاقة المتجددة مع فلاتر للسنة ومصدر الطاقة.",
                "desc_en": "A dashboard displaying a chart for renewable energy production with filters for year and energy source.",
                "charts": [
                    ("كمية الإنتاج المحلي من الطاقة الجديدة والمتجددة", "Local Production Quantity of New and Renewable Energy", "مخطط يعرض كميات الإنتاج المحلي لمصادر الطاقة المتجددة حسب السنة المحددة.", "A chart showing local production quantities of renewable energy sources by the selected year."),
                ],
            },
            {
                "name_ar": "الكميات المستوردة من النفط الخام ومشتقاته",
                "name_en": "Imported Quantities of Crude Oil and Derivatives",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لكميات استيراد النفط الخام ومشتقاته مع فلاتر للسنة ونوع المشتقات.",
                "desc_en": "A dashboard displaying a chart for imported crude oil and derivatives with filters for year and derivative type.",
                "charts": [
                    ("الكميات المستوردة من النفط الخام ومشتقاته", "Imported Quantities of Crude Oil and Derivatives", "مخطط يعرض كميات استيراد النفط الخام ومشتقاته حسب السنة ونوع المشتق المحدد.", "A chart showing imported quantities of crude oil and derivatives by year and selected derivative type."),
                ],
            },
            {
                "name_ar": "الطاقة الأولية المستهلكة",
                "name_en": "Primary Energy Consumed",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للطاقة الأولية المستهلكة مع فلاتر لنوع الطاقة والسنة.",
                "desc_en": "A dashboard displaying a chart for primary energy consumed with filters for energy type and year.",
                "charts": [
                    ("الطاقة الأولية المستهلكة", "Primary Energy Consumed", "مخطط يعرض كمية الطاقة الأولية المستهلكة حسب نوع الطاقة والسنة المحددة.", "A chart showing primary energy consumed by energy type and selected year."),
                ],
            },
            {
                "name_ar": "النتائج الرئيسية للمسح الصناعي",
                "name_en": "Main Results of the Industrial Survey",
                "desc_ar": "لوحة تعرض مخططات بيانية لنتائج المسح الصناعي مع فلاتر للنشاط الاقتصادي والسنة.",
                "desc_en": "A dashboard displaying charts for industrial survey results with filters for economic activity and year.",
                "charts": [
                    ("الإنتاج القائم", "Gross Output", "مخطط يعرض الإنتاج القائم للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing gross output for the selected economic activity by year."),
                    ("الاستهلاك الوسيط", "Intermediate Consumption", "مخطط يعرض الاستهلاك الوسيط للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing intermediate consumption for the selected economic activity by year."),
                    ("القيمة المضافة الإجمالية", "Gross Value Added", "مخطط يعرض القيمة المضافة الإجمالية للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing gross value added for the selected economic activity by year."),
                    ("الاهتلاك", "Depreciation", "مخطط يعرض الاهتلاك للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing depreciation for the selected economic activity by year."),
                    ("الضرائب على الإنتاج", "Taxes on Production", "مخطط يعرض الضرائب على الإنتاج للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing taxes on production for the selected economic activity by year."),
                    ("تكوين رأس المال الثابت الإجمالي", "Gross Fixed Capital Formation", "مخطط يعرض تكوين رأس المال الثابت الإجمالي للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing gross fixed capital formation for the selected economic activity by year."),
                    ("تعويضات العاملين", "Employee Compensation", "مخطط يعرض تعويضات العاملين للنشاط الاقتصادي المحدد حسب السنة.", "A chart showing employee compensation for the selected economic activity by year."),
                ],
            },
        ],
    },
    {
        "dashboard_ar": "الكهرباء",
        "dashboard_en": "Electricity",
        "tabs": [
            {
                "name_ar": "الكهرباء – معلومات عامة",
                "name_en": "Electricity – General Information",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لبيانات الطاقة الكهربائية مع فلاتر للبند والسنة.",
                "desc_en": "A dashboard displaying a chart for electricity data with filters for item and year.",
                "charts": [
                    ("بيانات عامة عن الطاقة الكهربائية", "General Electricity Data", "مخطط يعرض بيانات الطاقة الكهربائية للبند المحدد حسب السنة.", "A chart showing electricity data for the selected item by year."),
                ],
            },
            {
                "name_ar": "الكهرباء – تفصيلية",
                "name_en": "Electricity – Detailed",
                "desc_ar": "لوحة تعرض مخططات بيانية تفصيلية للطاقة الكهربائية مع فلتر للسنة.",
                "desc_en": "A dashboard displaying detailed electricity charts with a filter for year.",
                "charts": [
                    ("مصدر توليد الطاقة", "Power Generation Source", "مخطط يعرض إنتاج الطاقة حسب مصادر التوليد والسنة المحددة.", "A chart showing energy production by generation sources and selected year."),
                    ("الحمل الأقصى", "Maximum Load", "مخطط يعرض الحمل الأقصى للطاقة الكهربائية حسب السنة المحددة.", "A chart showing maximum electricity load by selected year."),
                    ("الصناعات الكبرى", "Major Industries", "مخطط يعرض بيانات الطاقة الكهربائية لقطاع الصناعات الكبرى حسب السنة المحددة.", "A chart showing electricity data for the major industries sector by selected year."),
                ],
            },
            {
                "name_ar": "استهلاك المملكة من الوقود",
                "name_en": "Kingdom Fuel Consumption",
                "desc_ar": "لوحة تعرض مخططات بيانية لاستهلاك الوقود والكهرباء مع فلاتر جانبية.",
                "desc_en": "A dashboard displaying charts for fuel and electricity consumption with side filters.",
                "charts": [
                    ("نصيب الفرد من الوقود", "Per Capita Fuel Consumption", "مخطط يعرض نصيب الفرد من استهلاك الوقود عبر السنوات.", "A chart showing per capita fuel consumption across years."),
                    ("استهلاك المملكة من الوقود", "Kingdom Fuel Consumption", "مخطط يعرض إجمالي استهلاك المملكة من الوقود عبر السنوات.", "A chart showing total kingdom fuel consumption across years."),
                    ("نسبة استهلاك قطاع الكهرباء من الاستهلاك الكلي", "Electricity Sector Share of Total Consumption", "مخطط يعرض نسبة استهلاك قطاع الكهرباء من إجمالي الاستهلاك عبر السنوات.", "A chart showing the electricity sector's share of total consumption across years."),
                    ("استهلاك قطاع الكهرباء", "Electricity Sector Consumption", "مخطط يعرض استهلاك قطاع الكهرباء من الوقود عبر السنوات.", "A chart showing electricity sector fuel consumption across years."),
                ],
            },
            {
                "name_ar": "أعداد المشتركين بالتيار الكهربائي حسب المصدر",
                "name_en": "Electricity Subscribers by Source",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لأعداد المشتركين بالتيار الكهربائي مع فلاتر للمصدر والسنة.",
                "desc_en": "A dashboard displaying a chart for electricity subscribers with filters for source and year.",
                "charts": [
                    ("أعداد المشتركين في التيار الكهربائي", "Electricity Subscribers", "مخطط يعرض أعداد المشتركين حسب مصدر التزويد والسنة المحددة.", "A chart showing subscriber counts by supply source and selected year."),
                ],
            },
            {
                "name_ar": "الطاقة الكهربائية حسب الاستخدام",
                "name_en": "Electricity by Use",
                "desc_ar": "لوحة تعرض مخططات بيانية للطاقة الكهربائية حسب الاستخدام مع فلاتر للسنة ونوع الاستخدام.",
                "desc_en": "A dashboard displaying charts for electricity by use with filters for year and use type.",
                "charts": [
                    ("كمية الطاقة الكهربائية", "Electricity Quantity", "مخطط يعرض كمية الطاقة الكهربائية حسب نوع الاستخدام والسنة المحددة.", "A chart showing electricity quantity by use type and selected year."),
                    ("النسبة إلى الإجمالي", "Share of Total", "مخطط يعرض نسبة كل نوع استخدام إلى إجمالي الاستهلاك حسب السنة المحددة.", "A chart showing each use type's share of total consumption by selected year."),
                ],
            },
        ],
    },
    {
        "dashboard_ar": "العمل والأجور",
        "dashboard_en": "Labor and Wages",
        "tabs": [
            {
                "name_ar": "حسب الجنس والمستوى التعليمي",
                "name_en": "By Gender and Educational Level",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للعاملين في القطاع العام مع فلاتر للسنة والمستوى التعليمي والجنس.",
                "desc_en": "A dashboard displaying a chart for public sector workers with filters for year, educational level, and gender.",
                "charts": [],
            },
            {
                "name_ar": "توزيع العاملين في القطاعين العام والخاص",
                "name_en": "Workers Distribution in Public and Private Sectors",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لتوزيع العاملين في القطاعين العام والخاص مع فلاتر للسنة والنشاط الاقتصادي والجنس.",
                "desc_en": "A dashboard displaying a chart for worker distribution in public and private sectors with filters for year, economic activity, and gender.",
                "charts": [],
            },
            {
                "name_ar": "حسب الإقليم والجنسية",
                "name_en": "By Region and Nationality",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للعاملين في منشآت القطاعين العام والخاص مع فلاتر للسنة والإقليم والجنسية.",
                "desc_en": "A dashboard displaying a chart for workers in public and private establishments with filters for year, region, and nationality.",
                "charts": [],
            },
            {
                "name_ar": "حسب النشاط الاقتصادي والجنس",
                "name_en": "By Economic Activity and Gender",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للعاملين في منشآت القطاعين العام والخاص مع فلاتر للسنة والنشاط الاقتصادي والجنس.",
                "desc_en": "A dashboard displaying a chart for workers in public and private establishments with filters for year, economic activity, and gender.",
                "charts": [],
            },
            {
                "name_ar": "حسب الجنسية والجنس",
                "name_en": "By Nationality and Gender",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للعاملين في منشآت القطاعين العام والخاص مع فلاتر للسنة والجنسية والجنس.",
                "desc_en": "A dashboard displaying a chart for workers in public and private establishments with filters for year, nationality, and gender.",
                "charts": [],
            },
            {
                "name_ar": "العاملون بأجر – منشآت القطاعين",
                "name_en": "Wage Earners – Public and Private Establishments",
                "desc_ar": "لوحة تعرض مخططات بيانية للعاملين بأجر وأجورهم وساعات عملهم مع فلاتر للسنة والمهنة والجنس.",
                "desc_en": "A dashboard displaying charts for wage earners, wages, and working hours with filters for year, occupation, and gender.",
                "charts": [],
            },
            {
                "name_ar": "العاملون بأجر – منشآت القطاع العام",
                "name_en": "Wage Earners – Public Sector Establishments",
                "desc_ar": "لوحة تعرض مخططات بيانية للعاملين بأجر في منشآت القطاع العام مع فلاتر للسنة والمهنة والجنس.",
                "desc_en": "A dashboard displaying charts for wage earners in public sector establishments with filters for year, occupation, and gender.",
                "charts": [],
            },
            {
                "name_ar": "العاملون بأجر – منشآت القطاع الخاص",
                "name_en": "Wage Earners – Private Sector Establishments",
                "desc_ar": "لوحة تعرض مخططات بيانية للعاملين بأجر في منشآت القطاع الخاص مع فلاتر للسنة والمهنة والجنس.",
                "desc_en": "A dashboard displaying charts for wage earners in private sector establishments with filters for year, occupation, and gender.",
                "charts": [],
            },
            {
                "name_ar": "عدد العاملين – القطاع العام",
                "name_en": "Number of Workers – Public Sector",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للعاملين في القطاع العام مع فلاتر للسنة والمهنة والجنسية.",
                "desc_en": "A dashboard displaying a chart for public sector workers with filters for year, occupation, and nationality.",
                "charts": [],
            },
            {
                "name_ar": "عدد العاملين – القطاع الخاص",
                "name_en": "Number of Workers – Private Sector",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لعدد العاملين في القطاع الخاص مع فلاتر للسنة والمهنة والجنسية.",
                "desc_en": "A dashboard displaying a chart for private sector workers with filters for year, occupation, and nationality.",
                "charts": [],
            },
        ],
    },
    {
        "dashboard_ar": "المعلومات والاتصالات",
        "dashboard_en": "Information and Communications",
        "tabs": [
            {
                "name_ar": "أعداد المشتركين في خدمات الهواتف الثابتة والخلوية",
                "name_en": "Fixed and Mobile Phone Service Subscribers",
                "desc_ar": "لوحة تعرض مخططات بيانية لأعداد المشتركين في خدمات الاتصالات مع فلتر للسنة.",
                "desc_en": "A dashboard displaying charts for communication service subscribers with a filter for year.",
                "charts": [],
            },
            {
                "name_ar": "توزيع الأسر حسب مؤشرات تكنولوجيا المعلومات",
                "name_en": "Household Distribution by ICT Indicators",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لتوزيع الأسر حسب مؤشرات تكنولوجيا المعلومات مع فلاتر للمؤشر والسنة.",
                "desc_en": "A dashboard displaying a chart for household distribution by ICT indicators with filters for indicator and year.",
                "charts": [],
            },
            {
                "name_ar": "التوزيع النسبي للأفراد حسب استخدام الحاسوب",
                "name_en": "Relative Distribution of Individuals by Computer Use",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للتوزيع النسبي لاستخدام الحاسوب مع فلاتر للسنة والجنس والعمر ونوع الاستخدام.",
                "desc_en": "A dashboard displaying a chart for relative computer use distribution with filters for year, gender, age, and use type.",
                "charts": [],
            },
            {
                "name_ar": "التوزيع النسبي للأفراد حسب استخدام الإنترنت",
                "name_en": "Relative Distribution of Individuals by Internet Use",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للتوزيع النسبي لاستخدام الإنترنت مع فلاتر للسنة والجنس والعمر ونوع الاستخدام.",
                "desc_en": "A dashboard displaying a chart for relative internet use distribution with filters for year, gender, age, and use type.",
                "charts": [
                    ("التوزيع النسبي للأفراد حسب استخدام الإنترنت", "Relative Distribution by Internet Use", "مخطط يعرض التوزيع النسبي لاستخدام الإنترنت حسب الجنس والفئة العمرية ونوع الاستخدام والسنة المحددة.", "A chart showing relative internet use distribution by gender, age group, use type, and selected year."),
                ],
            },
        ],
    },
    {
        "dashboard_ar": "حوادث الطرق",
        "dashboard_en": "Road Accidents",
        "tabs": [
            {
                "name_ar": "أخطاء السائقين المشتركين في حوادث الطرق",
                "name_en": "Driver Errors in Road Accidents",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لأخطاء السائقين في حوادث الطرق مع فلاتر للسنة ونوع الخطأ.",
                "desc_en": "A dashboard displaying a chart for driver errors in road accidents with filters for year and error type.",
                "charts": [],
            },
            {
                "name_ar": "المركبات المشتركة في حوادث الطرق",
                "name_en": "Vehicles Involved in Road Accidents",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للمركبات المشتركة في حوادث الطرق مع فلاتر للسنة والشهر ونوع المركبة.",
                "desc_en": "A dashboard displaying a chart for vehicles involved in road accidents with filters for year, month, and vehicle type.",
                "charts": [],
            },
            {
                "name_ar": "حوادث الطرق حسب نوع الحادث والمتضررين",
                "name_en": "Road Accidents by Accident Type and Affected Persons",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لحوادث الطرق مع فلاتر للسنة والمتضررين.",
                "desc_en": "A dashboard displaying a chart for road accidents with filters for year and affected persons.",
                "charts": [],
            },
            {
                "name_ar": "حوادث الطرق حسب المحافظة ونوع الحادث",
                "name_en": "Road Accidents by Governorate and Accident Type",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لحوادث الطرق مع فلاتر للسنة والمحافظة ونوع الحادث.",
                "desc_en": "A dashboard displaying a chart for road accidents with filters for year, governorate, and accident type.",
                "charts": [],
            },
            {
                "name_ar": "حوادث الطرق حسب الشهر ونوع الحادث",
                "name_en": "Road Accidents by Month and Accident Type",
                "desc_ar": "لوحة تعرض مخططاً بيانياً لحوادث الطرق مع فلاتر للسنة والشهر ونوع الحادث.",
                "desc_en": "A dashboard displaying a chart for road accidents with filters for year, month, and accident type.",
                "charts": [],
            },
            {
                "name_ar": "المصابين في حوادث الطرق حسب المصاب",
                "name_en": "Injured in Road Accidents by Injured Person Type",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للمصابين في حوادث الطرق مع فلاتر للسنة ونوع المصاب.",
                "desc_en": "A dashboard displaying a chart for injured persons in road accidents with filters for year and injured person type.",
                "charts": [],
            },
            {
                "name_ar": "المصابين في حوادث الطرق حسب درجة الإصابة",
                "name_en": "Injured in Road Accidents by Injury Degree",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للمصابين في حوادث الطرق مع فلاتر للسنة ودرجة الإصابة.",
                "desc_en": "A dashboard displaying a chart for injured persons in road accidents with filters for year and injury degree.",
                "charts": [],
            },
        ],
    },
    {
        "dashboard_ar": "الثقافة والإعلام",
        "dashboard_en": "Culture and Media",
        "tabs": [
            {
                "name_ar": "النشاطات الثقافية",
                "name_en": "Cultural Activities",
                "desc_ar": "لوحة تعرض مخططاً بيانياً للنشاطات الثقافية مع فلاتر لنوع النشاط والسنة.",
                "desc_en": "A dashboard displaying a chart for cultural activities with filters for activity type and year.",
                "charts": [],
            },
            {
                "name_ar": "المطابع ودور النشر والتوزيع",
                "name_en": "Printing Presses and Publishing Houses",
                "desc_ar": "لوحة تعرض مخططات بيانية للمطابع ودور النشر والتوزيع مع فلاتر للمحافظة والسنة.",
                "desc_en": "A dashboard displaying charts for printing presses and publishing houses with filters for governorate and year.",
                "charts": [],
            },
        ],
    },
]


def set_cell_shading(cell, color_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_paragraph_rtl(paragraph, rtl: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1" if rtl else "0")
    p_pr.append(bidi)


def add_styled_run(paragraph, text: str, bold: bool = False, size: int = 11, rtl: bool = False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if rtl:
        r_pr = run._r.get_or_add_rPr()
        rtl_el = OxmlElement("w:rtl")
        rtl_el.set(qn("w:val"), "1")
        r_pr.append(rtl_el)
    return run


def build_document() -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_styled_run(title, "أوصاف لوحات المعلومات", bold=True, size=18, rtl=True)
    title.add_run("\n")
    add_styled_run(title, "Dashboard Descriptions", bold=True, size=16)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_styled_run(subtitle, "وصف التبويبات والمخططات – عربي / إنجليزي", size=11, rtl=True, color=RGBColor(80, 80, 80))
    subtitle.add_run("\n")
    add_styled_run(subtitle, "Tab and Chart Descriptions – Arabic / English", size=11, color=RGBColor(80, 80, 80))

    doc.add_paragraph()

    for section in SECTIONS:
        heading = doc.add_paragraph()
        set_paragraph_rtl(heading, True)
        add_styled_run(heading, section["dashboard_ar"], bold=True, size=14, rtl=True, color=RGBColor(0, 102, 102))
        heading.add_run("  |  ")
        add_styled_run(heading, section["dashboard_en"], bold=True, size=14, color=RGBColor(0, 102, 102))

        for tab in section["tabs"]:
            tab_heading = doc.add_paragraph()
            set_paragraph_rtl(tab_heading, True)
            add_styled_run(tab_heading, "التبويب: ", bold=True, size=12, rtl=True)
            add_styled_run(tab_heading, tab["name_ar"], bold=True, size=12, rtl=True)
            tab_heading.add_run("\n")
            add_styled_run(tab_heading, "Tab: ", bold=True, size=12)
            add_styled_run(tab_heading, tab["name_en"], bold=True, size=12)

            table = doc.add_table(rows=2, cols=2)
            table.style = "Table Grid"
            table.autofit = False
            for col in table.columns:
                col.width = Inches(3.1)

            headers = [("وصف التبويب (عربي)", True), ("Tab Description (English)", False)]
            for i, (header, rtl) in enumerate(headers):
                cell = table.rows[0].cells[i]
                set_cell_shading(cell, "E6F2F2")
                p = cell.paragraphs[0]
                if rtl:
                    set_paragraph_rtl(p, True)
                add_styled_run(p, header, bold=True, size=10, rtl=rtl)

            desc_cells = [(tab["desc_ar"], True), (tab["desc_en"], False)]
            for i, (text, rtl) in enumerate(desc_cells):
                cell = table.rows[1].cells[i]
                p = cell.paragraphs[0]
                if rtl:
                    set_paragraph_rtl(p, True)
                add_styled_run(p, text, size=10, rtl=rtl)

            if tab["charts"]:
                doc.add_paragraph()
                chart_label = doc.add_paragraph()
                set_paragraph_rtl(chart_label, True)
                add_styled_run(chart_label, "وصف المخططات:", bold=True, size=11, rtl=True)
                chart_label.add_run("  ")
                add_styled_run(chart_label, "Chart Descriptions:", bold=True, size=11)

                chart_table = doc.add_table(rows=1, cols=4)
                chart_table.style = "Table Grid"
                chart_headers = [
                    ("اسم المخطط (عربي)", True),
                    ("Chart Name (English)", False),
                    ("الوصف (عربي)", True),
                    ("Description (English)", False),
                ]
                for i, (header, rtl) in enumerate(chart_headers):
                    cell = chart_table.rows[0].cells[i]
                    set_cell_shading(cell, "E6F2F2")
                    p = cell.paragraphs[0]
                    if rtl:
                        set_paragraph_rtl(p, True)
                    add_styled_run(p, header, bold=True, size=9, rtl=rtl)

                for chart in tab["charts"]:
                    row = chart_table.add_row().cells
                    values = [
                        (chart[0], True),
                        (chart[1], False),
                        (chart[2], True),
                        (chart[3], False),
                    ]
                    for i, (text, rtl) in enumerate(values):
                        p = row[i].paragraphs[0]
                        if rtl:
                            set_paragraph_rtl(p, True)
                        add_styled_run(p, text, size=9, rtl=rtl)

            doc.add_paragraph()

    return doc


if __name__ == "__main__":
    document = build_document()
    document.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
